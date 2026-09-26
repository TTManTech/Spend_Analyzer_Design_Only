"""Build data/spend-rates.json from data/Master_Data_V0.1.xlsx.

Spend rate for a month = total spend that month / total income that month,
rounded to a whole percent (half up). The same is done per calendar year for YTD.

- Spend: Transactions rows by calendar month of Date, excluding card payments
  (Spend / Income = "Payment"); merchant refunds are negative so they net against spend.
- Income: Income sheet rows by calendar month of Date.
- Categories: each month and year also lists its spend per Category of Spend
  (net of refunds, share of the period's spend) and that category's top 7
  purchases with each one's share of the category total. A merchant fills at
  most 3 of those rows; when it has more than 3 purchases, the ones not listed
  are summarised under "more" (count and total) so the page can show them as
  a small bubble.

Run: python3 scripts/build_spend_rates.py   (needs openpyxl)
"""
import json
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "Master_Data_V0.1.xlsx"
OUT = ROOT / "data" / "spend-rates.json"


def monthly_totals(ws, skip_kind=None):
    totals = defaultdict(Decimal)
    rows = ws.iter_rows(min_row=2, values_only=True)
    for _bank, date, _year, _desc, amount, _category, kind in rows:
        if date is None or amount is None or kind == skip_kind:
            continue
        totals[f"{date.year}-{date.month:02d}"] += Decimal(str(amount))
    return totals


TOP_N = 7
PER_MERCHANT = 3


def spend_rows(ws):
    rows = ws.iter_rows(min_row=2, values_only=True)
    for _bank, date, _year, desc, amount, category, kind in rows:
        if date is None or amount is None or kind == "Payment":
            continue
        yield date, desc, Decimal(str(amount)), category


def pct(part, whole):
    return float((part / whole * 100).quantize(Decimal("0.1"), ROUND_HALF_UP))


def category_breakdown(rows):
    """rows: (date, desc, amount, category) for one period."""
    totals, purchases = defaultdict(Decimal), defaultdict(list)
    for date, desc, amount, category in rows:
        totals[category] += amount
        if amount > 0:
            purchases[category].append((amount, date, desc))
    spend = sum((t for t in totals.values() if t > 0), Decimal(0))
    out = []
    for category, total in sorted(totals.items(), key=lambda kv: -kv[1]):
        if total <= 0:
            continue
        ranked = sorted(purchases[category], key=lambda p: (-p[0], p[1]))
        listed, per_merchant, counts = set(), defaultdict(int), defaultdict(int)
        for i, p in enumerate(ranked):
            counts[p[2]] += 1
            if len(listed) < TOP_N and per_merchant[p[2]] < PER_MERCHANT:
                listed.add(i)
                per_merchant[p[2]] += 1
        top = [p for i, p in enumerate(ranked) if i in listed]
        more = defaultdict(lambda: [0, Decimal(0)])
        for i, p in enumerate(ranked):
            if counts[p[2]] > PER_MERCHANT and i not in listed:
                more[p[2]][0] += 1
                more[p[2]][1] += p[0]
        out.append({
            "name": category,
            "total": float(total),
            "share": pct(total, spend),
            "top": [{"desc": d, "date": dt.strftime("%Y-%m-%d"), "amount": float(a), "share": pct(a, total)}
                    for a, dt, d in top],
            "more": [{"desc": d, "count": n, "amount": float(a), "share": pct(a, total)}
                     for d, (n, a) in sorted(more.items(), key=lambda kv: -kv[1][1])],
        })
    return out


def by_year(totals):
    years = defaultdict(Decimal)
    for key, amount in totals.items():
        years[key[:4]] += amount
    return years


def summarise(spend, income):
    out = {}
    for key in sorted(set(spend) | set(income)):
        s, i = spend.get(key, Decimal(0)), income.get(key, Decimal(0))
        rate = int((s / i * 100).quantize(Decimal(1), ROUND_HALF_UP)) if i else None
        out[key] = {"spend": float(s.quantize(Decimal("0.01"))), "income": float(i), "spendRate": rate}
    return out


def main():
    wb = load_workbook(SRC, data_only=True)
    spend = monthly_totals(wb["Transactions"], skip_kind="Payment")
    income = monthly_totals(wb["Income"])

    months = summarise(spend, income)
    years = summarise(by_year(spend), by_year(income))

    periods = defaultdict(list)
    for row in spend_rows(wb["Transactions"]):
        periods[f"{row[0].year}-{row[0].month:02d}"].append(row)
        periods[str(row[0].year)].append(row)
    for key, entry in list(months.items()) + list(years.items()):
        entry["categories"] = category_breakdown(periods.get(key, []))

    OUT.write_text(json.dumps({"source": SRC.name, "months": months, "years": years}, indent=2) + "\n")
    print(f"Wrote {OUT.relative_to(ROOT)} ({len(months)} months, {len(years)} years)")


if __name__ == "__main__":
    main()
